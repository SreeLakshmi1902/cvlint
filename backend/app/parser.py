import re
from docx import Document

BULLET_PATTERN = re.compile(
    r"^(?:"
    r"[^\w\s(\[{\"'\u201c\u2018$£€#]+"
    r"|\d{1,2}[.)]"
    r"|[a-z][.)]"
    r")\s+"   
)

def strip_bullet(text:str) -> tuple[str, bool]:
    match = BULLET_PATTERN.match(text)
    if match:
        return text[match.end():], True
    return text, False

def is_word_bullet(paragraph) -> bool:
    ppr = paragraph._p.pPr
    return ppr is not None and ppr.numPr is not None   
    
def extract_docx_text(file_path: str) -> str:
    document = Document(file_path)
    text = ""
    for paragraph in document.paragraphs:
        if paragraph.text.strip():
            text += paragraph.text + "\n"
    return text

def extract_docx_lines(files_path:str) -> list[dict]:
    document = Document(files_path)
    lines = []
    for paragraph in document.paragraphs:
        style = paragraph.style.name.lower() if paragraph.style is not None else ""
        word_bullet = is_word_bullet(paragraph)
        
        for raw in paragraph.text.split("\n"):
            text = re.sub(r"\s+", " ", raw).strip()
            if not text:
                continue
        
            text, text_bullet = strip_bullet(text)
        
            lines.append({
                "text": text,
                "is_bullet": word_bullet or text_bullet,
                "style": style,
            })
    return lines  

SECTION_ALIASES = {
    "summary": "summary", "professional summary": "summary", "profile": "summary",
    "about me": "summary", "objective": "summary", "career objective": "summary",
    "education": "education", "academic background": "education",
    "experience": "experience", "professional experience": "experience",
    "work experience": "experience", "employment history": "experience",
    "work history": "experience",
    "skills": "skills", "technical skills": "skills", "core skills": "skills",
    "key skills": "skills", "technical expertise": "skills",
    "projects": "projects", "personal projects": "projects",
    "academic projects": "projects", "key projects": "projects",
    "certifications": "certifications", "certificates": "certifications",
    "licenses and certifications": "certifications",
    "awards": "awards", "achievements": "awards",
    "publications": "publications",
}

SUMMARY_MIN_WORDS = 15


def normalize_heading(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip().rstrip(":").lower()


def looks_like_contact(text: str) -> bool:
    lowered = text.lower()
    has_phone = re.search(r"\d{3}.?\d{3}.?\d{3}", text) is not None
    return "@" in text or "linkedin" in lowered or "github" in lowered or has_phone


def move_summary_out_of_header(sections: dict[str, list[dict]]) -> None:
    keep, summary = [], []
    for position, line in enumerate(sections["header"]):
        is_long_prose = len(line["text"].split()) >= SUMMARY_MIN_WORDS
        if position > 0 and is_long_prose and not looks_like_contact(line["text"]):
            summary.append(line)
        else:
            keep.append(line)

    sections["header"] = keep
    if summary:
        sections["summary"] = summary + sections.get("summary", [])


def build_sections(lines: list[dict]) -> dict[str, list[dict]]:
    sections = {"header": []}
    current = "header"

    for index, line in enumerate(lines):
        key = normalize_heading(line["text"])
        canonical = SECTION_ALIASES.get(key)
        heading_style = line["style"].startswith("heading") and index > 0

        if not line["is_bullet"] and (canonical or heading_style):
            current = canonical or key
            sections.setdefault(current, [])
            continue

        sections[current].append(line)

    move_summary_out_of_header(sections)
    return sections  

SKILL_SEPARATORS = re.compile(r"[,;|\u2022]\s*(?![^()]*\))")
DESCRIPTION_DASH = re.compile(r"\s[--]\s")


def is_category_label(line: dict) -> bool:
    text = line["text"]
    return not line["is_bullet"] and ":" not in text and not SKILL_SEPARATORS.search(text)


def split_skill_line(text: str) -> list[str]:
    if ":" in text:
        text = text.split(":", 1)[1]
    items = []
    for item in SKILL_SEPARATORS.split(text):
        item = DESCRIPTION_DASH.split(item)[0]
        item = item.strip().rstrip(".").strip()
        if item:
            items.append(item)
    return items


def extract_skill_list(sections: dict[str, list[dict]]) -> list[str]:
    skills, seen = [], set()
    for line in sections.get("skills", []):
        if is_category_label(line):
            continue
        for skill in split_skill_line(line["text"]):
            if skill.lower() not in seen:
                seen.add(skill.lower())
                skills.append(skill)
    return skills

def extract_sections(text: str):
    section_headers = {
        "education", 
        "professional experience", 
        "technical skills", 
        "projects",
        "certifications"
    }
    sections = {}
    current_section = "header"
    sections[current_section] = []
    for line in text.split("\n"):
        line = line.strip()
        if not line:
            continue
        normalized_line = line.lower()
        if normalized_line in section_headers:
            current_section = normalized_line
            sections[current_section] = []
        else:
            sections[current_section].append(line)
    return sections

def extract_skills(sections):
    skills_text = ""
    if "technical skills" in sections:
        skills_text = " ".join(sections["technical skills"])
    skills = [skill.strip() for skill in skills_text.split(",") if skill.strip()]
    return skills